"""Pipeline tests. Run from the repo root:  python3 -m unittest discover -s pipeline/tests -v

Standard library + openpyxl only. The model step is never called for real: it's replaced by a stub.
"""
import datetime
import io
import json
import re
import subprocess
import sys
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "pipeline"))
import dockparse  # noqa: E402
from dockparse import _merge, model  # noqa: E402

SAMPLE = ROOT / "sample_data" / "Dock Schedule - Synthetic Sample.xlsx"
ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")
ISSUE_CODES = {"NO_BERTH", "OUTSIDE_MONTH_COLUMNS", "UNPARSEABLE_CELL", "HEADER_YEAR_MISMATCH", "DUPLICATE_CARRYOVER",
               "ANNOTATION_SKIPPED", "UNKNOWN_FORMAT", "TEMPLATE_BAD_HEADER", "INVALID_VALUE", "DUPLICATE_NAME",
               "MODEL_CLASSIFIED"}


def check_shape(t, d):
    """The same shape shared/pipeline.ts (ParsedWorkbookSchema) enforces on the TS side."""
    t.assertEqual(d["version"], 1)
    t.assertIn(d["format"], ("template", "legacy_grid", None))
    t.assertEqual(set(d["stats"]), {"sheets", "cells", "modelCalls"})
    for b in d["berths"]:
        t.assertEqual(set(b), {"name", "kind", "lengthFt", "sortOrder"})
        t.assertIn(b["kind"], ("berth", "section"))
        t.assertTrue(b["lengthFt"] is None or b["lengthFt"] > 0)
    for v in d["vessels"]:
        t.assertEqual(set(v), {"name", "lengthFt", "draftFt", "operator", "notes"})
    rows = d["rows"] + [i["row"] for i in d["issues"] if i["row"]]
    for r in rows:
        t.assertEqual(set(r), {"berthLabel", "occupantType", "title", "startDate", "endDate", "notes", "sheet", "cell",
                               "classifiedBy"})
        t.assertIn(r["occupantType"], ("vessel", "event", "closure"))
        t.assertIn(r["classifiedBy"], ("regex", "template", "model"))
        t.assertRegex(r["startDate"], ISO)
        t.assertRegex(r["endDate"], ISO)
        t.assertTrue(r["title"])
    for i in d["issues"]:
        t.assertIn(i["code"], ISSUE_CODES)
        t.assertIn(i["severity"], ("error", "warning", "info"))
    json.dumps(d)  # serialisable


def xlsx(build) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class SampleWorkbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with mock.patch.object(model, "label", return_value=({}, 0)):
            cls.d = dockparse.run(SAMPLE.read_bytes())

    def test_shape(self):
        check_shape(self, self.d)

    def test_berths_match_the_default_fleet(self):
        got = [(b["name"], b["kind"], b["lengthFt"]) for b in self.d["berths"]]
        defaults = json.loads((ROOT / "backend" / "seed" / "defaults.json").read_text())["berths"]
        self.assertEqual(got, [(b["name"], b["kind"], b["lengthFt"]) for b in defaults])

    def test_golden_counts(self):
        from collections import Counter
        codes = Counter(i["code"] for i in self.d["issues"])
        self.assertEqual(self.d["format"], "legacy_grid")
        self.assertEqual(self.d["stats"]["sheets"], 23)
        self.assertTrue(1900 <= len(self.d["rows"]) <= 2300, len(self.d["rows"]))
        self.assertTrue(200 <= codes["NO_BERTH"] <= 320, codes["NO_BERTH"])
        self.assertEqual(codes["DUPLICATE_CARRYOVER"], 3)
        self.assertEqual(codes["HEADER_YEAR_MISMATCH"], 2)
        self.assertTrue(90 <= codes["ANNOTATION_SKIPPED"] <= 120)
        self.assertTrue(80 <= codes["OUTSIDE_MONTH_COLUMNS"] <= 100)

    def test_2010_sheet_years_inferred_not_trusted(self):
        # the 2010 sheet labels its last blocks "NOVEMBER 2018" / "DECEMBER 2018"
        on_2010 = [r for r in self.d["rows"] if r["sheet"] == "2010"]
        self.assertTrue(on_2010)
        self.assertTrue(all(r["startDate"].startswith("2010") for r in on_2010))

    def test_registry_lengths_attached(self):
        with_len = [v for v in self.d["vessels"] if v["lengthFt"] is not None]
        self.assertEqual(len(with_len), 164)          # the registry size (defaults.json)
        self.assertLessEqual(max(v["lengthFt"] for v in with_len), 410)

    def test_no_duplicate_stays(self):
        keys = [(r["berthLabel"], r["title"].upper(), r["startDate"]) for r in self.d["rows"]]
        self.assertEqual(len(keys), len(set(keys)))


class ModelStep(unittest.TestCase):
    def test_model_labels_only_leftovers_and_is_flagged(self):
        seen = {}

        def fake(texts, api_key=None):
            seen["texts"] = list(texts)
            return {"Bunker barge": "vessel"}, 1
        with mock.patch.object(model, "label", side_effect=fake):
            d = dockparse.run(SAMPLE.read_bytes())
        self.assertEqual(seen["texts"], ["Bunker barge"])     # regex placed everything else
        self.assertEqual(d["stats"]["modelCalls"], 1)
        by_model = [r for r in d["rows"] if r["classifiedBy"] == "model"]
        self.assertTrue(by_model and all(r["occupantType"] == "vessel" for r in by_model))
        self.assertFalse(any(i["code"] == "UNPARSEABLE_CELL" for i in d["issues"]))
        self.assertTrue(any(i["code"] == "MODEL_CLASSIFIED" for i in d["issues"]))

    def test_no_key_means_no_call(self):
        with mock.patch.dict("os.environ", {}, clear=True):
            self.assertEqual(model.label(["anything"]), ({}, 0))


class Template(unittest.TestCase):
    def build(self, wb):
        b = wb.create_sheet("Berths")
        b.append(["Name", "Kind", "Length (ft)", "Order"])
        b.append(["e.g. Example Pier", "berth", 100, 1])
        b.append(["Main Pier", "berth", 300, 1])
        b.append(["Slips", "section", None, 2])
        b.append(["Bad", "dock", 10, 3])
        b.append(["main pier", "berth", 200, 4])
        b.append(["No Length", "berth", None, 5])
        v = wb.create_sheet("Vessels")
        v.append(["Name", "Length (ft)", "Draft (ft)", "Operator", "Notes"])
        v.append(["R/V SEA LION", "120'", 9, "WHOI", None])
        v.append(["Unknown Len", None, None, None, None])
        v.append(["Negative", -5, None, None, None])
        k = wb.create_sheet("Bookings")
        k.append(["Berth", "Type", "Vessel / Title", "Start", "End", "Notes"])
        k.append(["Main Pier", "vessel", "r/v sea lion", datetime.datetime(2027, 3, 5), datetime.datetime(2027, 3, 9), None])
        k.append(["Slips", "event", "Sail day", "2027-04-01", "2027-04-01", "bring flags"])
        k.append(["Main Pier", "vessel", "M/V Newcomer", "2027-05-01", "2027-05-03", None])
        k.append(["Main Pier", "party", "x", "2027-05-01", "2027-05-01", None])
        k.append(["Main Pier", "closure", "Crane", "2027-02-30", "2027-03-01", None])
        k.append([None, "event", "Nowhere", "2027-06-01", "2027-06-01", None])
        k.append(["Main Pier", "event", "Backwards", "2027-07-09", "2027-07-01", None])   # rules.ts judges this

    def setUp(self):
        self.d = dockparse.run(xlsx(self.build))

    def test_shape_and_format(self):
        check_shape(self, self.d)
        self.assertEqual(self.d["format"], "template")

    def test_berths(self):
        self.assertEqual([(b["name"], b["kind"], b["lengthFt"]) for b in self.d["berths"]],
                         [("Main Pier", "berth", 300), ("Slips", "section", None)])

    def test_vessels_incl_unlisted_booked_one(self):
        names = {v["name"]: v["lengthFt"] for v in self.d["vessels"]}
        self.assertEqual(names, {"R/V Sea Lion": 120, "Unknown Len": None, "M/V Newcomer": None})

    def test_rows(self):
        rows = {r["title"]: r for r in self.d["rows"]}
        self.assertEqual(set(rows), {"R/V Sea Lion", "Sail day", "M/V Newcomer", "Backwards"})
        self.assertEqual((rows["R/V Sea Lion"]["startDate"], rows["R/V Sea Lion"]["endDate"]), ("2027-03-05", "2027-03-09"))
        self.assertEqual(rows["Sail day"]["notes"], "bring flags")
        self.assertTrue(all(r["classifiedBy"] == "template" for r in rows.values()))

    def test_issues(self):
        got = sorted((i["code"], i["cell"]) for i in self.d["issues"])
        self.assertEqual(got, sorted([
            ("INVALID_VALUE", "B5"),     # kind "dock"
            ("DUPLICATE_NAME", "A6"),    # "main pier" again
            ("INVALID_VALUE", "C7"),     # berth without length
            ("INVALID_VALUE", "B4"),     # negative vessel length
            ("INVALID_VALUE", "B5"),     # type "party"  (Bookings!B5)
            ("INVALID_VALUE", "D6"),     # 2027-02-30
            ("NO_BERTH", "A7"),          # empty berth → issue with row
        ]))
        no_berth = next(i for i in self.d["issues"] if i["code"] == "NO_BERTH")
        self.assertEqual(no_berth["row"]["title"], "Nowhere")

    def test_bad_header_skips_only_that_sheet(self):
        def build(wb):
            self.build(wb)
            ws = wb["Vessels"]
            ws["B1"] = "LOA"
        d = dockparse.run(xlsx(build))
        self.assertEqual([i["sheet"] for i in d["issues"] if i["code"] == "TEMPLATE_BAD_HEADER"], ["Vessels"])
        self.assertEqual(len(d["berths"]), 2)


class Edges(unittest.TestCase):
    def test_unknown_format(self):
        d = dockparse.run(xlsx(lambda wb: wb.create_sheet("Stuff").append(["hello"])))
        check_shape(self, d)
        self.assertIsNone(d["format"])
        self.assertEqual([i["code"] for i in d["issues"]], ["UNKNOWN_FORMAT"])

    def test_not_an_xlsx(self):
        d = dockparse.run(b"definitely not a zip")
        self.assertEqual([i["code"] for i in d["issues"]], ["UNKNOWN_FORMAT"])

    def test_merge_joins_across_month_boundary_but_not_gaps(self):
        r = lambda s, e: {"berthLabel": "B", "occupantType": "vessel", "title": "R/V A", "startDate": s, "endDate": e,
                          "notes": None, "sheet": "2020", "cell": "B1", "classifiedBy": "regex"}
        out = _merge([r("2020-01-29", "2020-01-31"), r("2020-02-01", "2020-02-03"), r("2020-02-05", "2020-02-05")])
        self.assertEqual([(x["startDate"], x["endDate"]) for x in out],
                         [("2020-01-29", "2020-02-03"), ("2020-02-05", "2020-02-05")])

    def test_cli_stdin_roundtrip(self):
        p = subprocess.run([sys.executable, "-m", "pipeline.cli", "--stdin", "--no-model"], cwd=ROOT,
                           input=xlsx(lambda wb: wb.create_sheet("Berths").append(["Name", "Kind", "Length (ft)", "Order"])),
                           capture_output=True, check=True)
        self.assertEqual(json.loads(p.stdout)["format"], "template")


if __name__ == "__main__":
    unittest.main()
