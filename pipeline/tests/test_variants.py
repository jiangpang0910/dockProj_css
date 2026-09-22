"""Four workbooks that all say the same thing, written four ways (scripts/make_variants.py), each shipped with the
truth it must yield. A reader is right when its output EQUALS the truth: no hand-written expectations.

    python3 -m unittest discover -s pipeline/tests -v
"""
import datetime as dt
import io
import json
import sys
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "pipeline"))
import dockparse  # noqa: E402
from dockparse import embed, model  # noqa: E402
from dockparse.audit import audit  # noqa: E402
from dockparse.classify import vessel_key  # noqa: E402
from dockparse.table import find_header, header_scores, map_columns  # noqa: E402
from dockparse.values import berth_label, parse_date, parse_range, split_name_length  # noqa: E402
from test_pipeline import check_shape, xlsx  # noqa: E402

VARIANTS = ROOT / "sample_data" / "variants"
FORMAT = {"1-grid-clone": "legacy_grid", "2-one-tab": "table", "3-renamed-columns": "table", "4-messy": "table"}


class Variant:
    """One workbook against its truth. Subclasses name the file."""
    NAME = ""

    @classmethod
    def setUpClass(cls):
        cls.data = (VARIANTS / f"{cls.NAME}.xlsx").read_bytes()
        cls.truth = json.loads((VARIANTS / f"{cls.NAME}.truth.json").read_text())
        cls.out = dockparse.run(cls.data, use_model=False)

    def test_shape_and_format(self):
        check_shape(self, self.out)
        self.assertEqual(self.out["format"], FORMAT[self.NAME])

    def test_every_stay_and_nothing_else(self):
        want = {(b["berth"], vessel_key(b["vessel"]), b["start"], b["end"]) for b in self.truth["bookings"]}
        got = {(r["berthLabel"], vessel_key(r["title"]), r["startDate"], r["endDate"]) for r in self.out["rows"]}
        self.assertEqual(want - got, set(), "stays the file states but the reader missed")
        self.assertEqual(got - want, set(), "stays the reader invented")
        self.assertTrue(all(r["occupantType"] == "vessel" for r in self.out["rows"]))

    def test_deliberate_conflicts_are_read_not_dropped(self):
        # the parser reports what the file says; whether it's allowed is rules.ts's job
        got = {(r["berthLabel"], vessel_key(r["title"]), r["startDate"]) for r in self.out["rows"]}
        clashes = [b for b in self.truth["bookings"] if b["conflict"]]
        self.assertEqual(len(clashes), 10)
        for b in clashes:
            self.assertIn((b["berth"], vessel_key(b["vessel"]), b["start"]), got)

    def test_vessel_lengths_stated_are_recovered_and_none_invented(self):
        got = {v["name"]: v for v in self.out["vessels"]}
        self.assertEqual(set(got), {v["name"] for v in self.truth["vessels"]})
        for v in self.truth["vessels"]:
            g = got[v["name"]]
            if v["lengthFt"] is not None:
                self.assertEqual(g["lengthFt"], v["lengthFt"], v["name"])
            elif g["lengthFt"] is not None:   # only a namesake's length may fill a blank, and it must say so
                self.assertIn("same name", g["notes"] or "", v["name"])

    def test_berths(self):
        want = {(b["name"], b["kind"], b["lengthFt"]) for b in self.truth["berths"]}
        got = {(b["name"], b["kind"], b["lengthFt"]) for b in self.out["berths"]}
        self.assertEqual(got, want)

    def test_no_errors_or_warnings(self):
        self.assertEqual([i for i in self.out["issues"] if i["severity"] != "info"], [])

    def test_every_cell_is_accounted_for(self):
        self.assertEqual(audit(self.data)["unaccounted"], [])

    def test_the_mapping_is_reported(self):
        if FORMAT[self.NAME] == "table":
            sheets = {i["sheet"] for i in self.out["issues"] if i["code"] == "COLUMN_MAPPING"}
            self.assertEqual(sheets, {r["sheet"] for r in self.out["rows"]} | sheets)
            self.assertTrue(sheets)


class GridClone(Variant, unittest.TestCase):
    NAME = "1-grid-clone"


class OneTab(Variant, unittest.TestCase):
    NAME = "2-one-tab"


class RenamedColumns(Variant, unittest.TestCase):
    NAME = "3-renamed-columns"


class Messy(Variant, unittest.TestCase):
    NAME = "4-messy"


class Values(unittest.TestCase):
    def test_dates_in_any_spelling(self):
        d = dt.date(2016, 3, 3)
        for s in ["2016-03-03", "3 Mar 2016", "Mar 3, 2016", "3rd March 16", "03/03/2016", "3.3.2016", "2016/03/03 09:00"]:
            self.assertEqual(parse_date(s), d, s)
        self.assertEqual(parse_date("03/05/2015", "DMY"), dt.date(2015, 5, 3))
        self.assertEqual(parse_date("03/05/2015", "MDY"), dt.date(2015, 3, 5))
        self.assertEqual(parse_date(42433), dt.date(2016, 3, 4))          # an Excel serial
        self.assertEqual(parse_date(dt.datetime(2016, 3, 3, 9)), d)
        for s in ["hello", "12 Jan", "2016", "", None, True]:
            self.assertIsNone(parse_date(s), s)

    def test_ranges_borrow_the_missing_parts(self):
        a, b = dt.date(2016, 3, 3), dt.date(2016, 3, 7)
        for s in ["3 Mar - 7 Mar 2016", "3–7 Mar 2016", "Mar 3-7, 2016", "2016-03-03 to 2016-03-07", "3/3/16-3/7/16"]:
            self.assertEqual(parse_range(s), (a, b), s)
        self.assertEqual(parse_range("28 Dec - 2 Jan 2017"), (dt.date(2016, 12, 28), dt.date(2017, 1, 2)))
        self.assertIsNone(parse_range("2016-03-03"))
        self.assertIsNone(parse_range("Mar 3 - Mar 7"))                     # no year anywhere

    def test_name_and_length_in_one_cell(self):
        self.assertEqual(split_name_length("North Bridge Speedster 46'"), ("North Bridge Speedster", 46, False))
        self.assertEqual(split_name_length("North Bridge Speedster 46 ft"), ("North Bridge Speedster", 46, False))
        self.assertEqual(split_name_length("F/V Blue Tern (85')"), ("F/V Blue Tern", 85, False))
        self.assertEqual(split_name_length("R/V Deep Cove 85 (LOA 92)"), ("R/V Deep Cove", 92, False))
        # a bare number is a length only when the column is known to carry lengths; otherwise it's flagged
        self.assertEqual(split_name_length("North Bridge Speedster 46"), ("North Bridge Speedster 46", None, True))
        self.assertEqual(split_name_length("North Bridge Speedster 46", bare_ok=True), ("North Bridge Speedster", 46, False))
        for name in ["M/Y Sea Fox II", "Sea Fox 2", "Barge 42"]:
            self.assertEqual(split_name_length(name, bare_ok=True), (name, None, False), name)

    def test_berth_labels(self):
        self.assertEqual(berth_label("NORTH PIER WEST (410 ft)"), ("North Pier West", 410))
        self.assertEqual(berth_label("North Pier West - 410'"), ("North Pier West", 410))
        self.assertEqual(berth_label("Small Craft Slips"), ("Small Craft Slips", None))


class TableReader(unittest.TestCase):
    def test_headers_by_wording(self):
        self.assertEqual(max(header_scores("Boat"), key=header_scores("Boat").get), "vessel")
        self.assertEqual(max(header_scores("Dock"), key=header_scores("Dock").get), "berth")
        self.assertEqual(header_scores("Arrival date")["start"], 1.0)
        self.assertEqual(header_scores("VESSEL (LOA)")["vessel"], 1.0)
        self.assertEqual(header_scores(None), {})

    def test_values_decide_when_headers_dont(self):
        headers = {1: "Col A", 2: "Col B", 3: "Col C", 4: "Col D"}
        cols = {1: ["R/V Kestrel", "F/V Osprey", "R/V Kestrel"], 2: ["North Pier", "South Float", "North Pier"],
                3: [dt.date(2027, 3, 1), dt.date(2027, 3, 4), dt.date(2027, 3, 9)],
                4: [dt.date(2027, 3, 3), dt.date(2027, 3, 6), dt.date(2027, 3, 10)]}
        mapping, _ = map_columns(headers, cols, set(), set())
        self.assertEqual(mapping, {1: "vessel", 2: "berth", 3: "start", 4: "end"})

    def test_a_row_with_lengths_is_never_the_header(self):
        rows = {1: {1: "berth lengths"}, 2: {1: "North Pier West", 2: "410 ft"}, 3: {1: "Inner Channel", 2: "90 ft"}}
        self.assertEqual(find_header(rows, set(), set()), (None, {}))

    def test_length_in_the_vessel_cell_of_a_booking(self):
        def build(wb):
            ws = wb.create_sheet("Log")
            ws.append(["Boat", "Dock", "In", "Out"])
            ws.append(["North Bridge Speedster 46'", "Inner Channel (90 ft)", "2027-03-01", "2027-03-03"])
            ws.append(["North Bridge Speedster 46'", "Inner Channel (90 ft)", "2027-03-10", "2027-03-12"])
            ws.append(["Sea Fox II", "Inner Channel (90 ft)", "2027-03-20", "2027-03-21"])
        d = dockparse.run(xlsx(build), use_model=False)
        check_shape(self, d)
        self.assertEqual(d["format"], "table")
        self.assertEqual({v["name"]: v["lengthFt"] for v in d["vessels"]}, {"North Bridge Speedster": 46, "Sea Fox II": None})
        self.assertEqual([(b["name"], b["kind"], b["lengthFt"]) for b in d["berths"]], [("Inner Channel", "berth", 90)])
        self.assertEqual([r["title"] for r in d["rows"]], ["North Bridge Speedster", "North Bridge Speedster", "Sea Fox II"])
        self.assertEqual([i["code"] for i in d["issues"]], ["COLUMN_MAPPING"])

    def test_csv(self):
        text = "Ship,Dock,Arrive,Depart\nR/V Test,North Pier West,2027-03-10,2027-03-13\nF/V Two,North Pier West,3 Apr 2027,5 Apr 2027\n"
        d = dockparse.run(text.encode(), use_model=False)
        self.assertEqual(d["format"], "table")
        self.assertEqual([(r["title"], r["berthLabel"], r["startDate"], r["endDate"]) for r in d["rows"]],
                         [("R/V Test", "North Pier West", "2027-03-10", "2027-03-13"), ("F/V Two", "North Pier West", "2027-04-03", "2027-04-05")])

    def test_unmapped_column_and_ambiguous_dates_are_reported(self):
        def build(wb):
            ws = wb.create_sheet("S")
            ws.append(["Boat", "Dock", "In", "Out", "Skipper"])
            ws.append(["R/V A", "North Pier", "03/05/2027", "04/05/2027", "Sam"])
            ws.append(["R/V B", "North Pier", "06/05/2027", "07/05/2027", "Alex"])
        d = dockparse.run(xlsx(build), use_model=False)
        codes = {i["code"] for i in d["issues"]}
        self.assertIn("UNMAPPED_COLUMN", codes)                     # "Skipper" means nothing to us; said so
        self.assertIn("AMBIGUOUS_VALUE", codes)                     # 03/05 read as month/day, nothing proved otherwise
        self.assertEqual(d["rows"][0]["startDate"], "2027-03-05")

    def test_a_day_above_twelve_settles_the_order_for_the_column(self):
        def build(wb):
            ws = wb.create_sheet("S")
            ws.append(["Boat", "Dock", "In", "Out"])
            ws.append(["R/V A", "North Pier", "03/05/2027", "04/05/2027"])
            ws.append(["R/V B", "North Pier", "14/05/2027", "15/05/2027"])
        d = dockparse.run(xlsx(build), use_model=False)
        self.assertEqual([r["startDate"] for r in d["rows"]], ["2027-05-03", "2027-05-14"])
        self.assertNotIn("AMBIGUOUS_VALUE", {i["code"] for i in d["issues"]})

    @unittest.skipUnless(embed.available(), "models/minilm not present")
    def test_headers_matched_by_meaning_with_the_local_model(self):
        def build(wb):
            ws = wb.create_sheet("Log")
            ws.append(["Watercraft", "Tied up", "Came alongside", "Departed", "Skipper"])   # none is a synonym we list
            ws.append(["1042", "7", dt.date(2027, 3, 1), dt.date(2027, 3, 2), "Sam"])       # hull and berth NUMBERS
            ws.append(["1043", "9", dt.date(2027, 3, 3), dt.date(2027, 3, 4), "Alex"])
            ws.append(["1044", "7", dt.date(2027, 3, 5), dt.date(2027, 3, 6), "Kim"])
        d = dockparse.run(xlsx(build), use_model=False)
        self.assertEqual(d["format"], "table")
        self.assertEqual([(r["title"], r["berthLabel"], r["startDate"], r["endDate"]) for r in d["rows"]],
                         [("1042", "7", "2027-03-01", "2027-03-02"), ("1043", "9", "2027-03-03", "2027-03-04"), ("1044", "7", "2027-03-05", "2027-03-06")])
        mapping = next(i["message"] for i in d["issues"] if i["code"] == "COLUMN_MAPPING")
        self.assertIn("A (Watercraft) → vessel [by meaning]", mapping)
        self.assertIn("B (Tied up) → berth [by meaning]", mapping)
        self.assertNotIn("Skipper", mapping)                                   # a person: decoy wins, column left out
        self.assertEqual([i["message"][:24] for i in d["issues"] if i["code"] == "UNMAPPED_COLUMN"], ['Column E ("Skipper", 3 c'])

    def test_without_the_model_nothing_changes(self):
        with mock.patch.object(embed, "available", return_value=False):
            self.assertEqual(header_scores("Watercraft"), {})
            self.assertEqual(header_scores("Boat"), {"vessel": 1.0})

    def test_a_column_without_dates_is_never_a_date_column(self):
        headers = {1: "Boat", 2: "Departed"}                                    # "Departed" hints end, but the values are names
        cols = {1: ["R/V A", "R/V B", "R/V C"], 2: ["Sam", "Alex", "Kim"]}
        mapping, _ = map_columns(headers, cols, set(), set())
        self.assertEqual(mapping, {1: "vessel"})

    def test_model_proposes_a_layout_only_when_nothing_matched(self):
        def build(wb):
            ws = wb.create_sheet("Hulls")
            for r in [["1042", 7, "2027-03-01", "2027-03-02"], ["1043", 9, "2027-03-03", "2027-03-04"],
                      ["1044", 7, "2027-03-05", "2027-03-06"]]:
                ws.append(r)                                        # hull and berth NUMBERS: nothing says "vessel"
        spec = {"Hulls": {"headerRow": None, "columns": {1: "vessel", 2: "berth", 3: "start", 4: "end"}}}
        with mock.patch.object(model, "propose_layout", return_value=(spec, 1)) as p:
            d = dockparse.run(xlsx(build), use_model=True)
        p.assert_called_once()
        self.assertEqual(d["format"], "table")
        self.assertEqual([(r["title"], r["berthLabel"]) for r in d["rows"]], [("1042", "7"), ("1043", "9"), ("1044", "7")])
        self.assertIn("MODEL_LAYOUT", {i["code"] for i in d["issues"]})
        self.assertEqual(d["stats"]["modelCalls"], 1)
        with mock.patch.object(model, "propose_layout", return_value=({}, 1)):
            d = dockparse.run(xlsx(build), use_model=True)          # the model had nothing: still no guess
        self.assertIsNone(d["format"])
        self.assertEqual([i["code"] for i in d["issues"]], ["UNKNOWN_FORMAT"])

    def test_garbage_from_the_model_is_dropped(self):
        with mock.patch.object(model, "_ask", return_value=({"X": {"headerRow": 1, "columns": {"A": "vessel", "B": "banana", "ZZ": "start"}},
                                                             "Y": "nonsense"}, 1)):
            specs, calls = model.propose_layout([{"title": "X", "rows": [], "fields": []}], api_key="k")
        self.assertEqual(specs, {"X": {"headerRow": 1, "columns": {1: "vessel", 702: "start"}}})
        self.assertEqual(calls, 1)


if __name__ == "__main__":
    unittest.main()
