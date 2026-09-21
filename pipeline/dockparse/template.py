"""Our upload template (public/dock-template.xlsx): plain tables, one header row, one record per row.

Sheets (all optional, names case-insensitive):
  Berths   : Name | Kind | Length (ft) | Order
  Vessels  : Name | Length (ft) | Draft (ft) | Operator | Notes
  Bookings : Berth | Type | Vessel / Title | Start | End | Notes
A row whose first cell starts with "e.g." is the template's example and is skipped. Python only reads and shapes
the data; whether a booking is allowed is decided later by rules.ts.
"""
import datetime
import re

from openpyxl.utils import get_column_letter

from .classify import clean, vessel_display, vessel_key

HEADERS = {
    "berths": ["name", "kind", "length (ft)", "order"],
    "vessels": ["name", "length (ft)", "draft (ft)", "operator", "notes"],
    "bookings": ["berth", "type", "vessel / title", "start", "end", "notes"],
}
ISO_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
FEET_RE = re.compile(r"^\s*(\d+(?:\.\d+)?)\s*(?:ft|feet|['’′])?\s*$", re.I)


def is_template(wb) -> bool:
    return any(t.strip().lower() in HEADERS for t in wb.sheetnames)


def _sheet(wb, name):
    for t in wb.sheetnames:
        if t.strip().lower() == name:
            return wb[t]
    return None


def _text(v):
    if v is None:
        return None
    s = clean(v)
    return s or None


def _feet(v):
    """→ (value | None, ok). Blank is (None, True)."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, True
    if isinstance(v, (int, float)) and v > 0:
        return (int(v) if float(v).is_integer() else float(v)), True
    if isinstance(v, str) and (m := FEET_RE.match(v)) and float(m[1]) > 0:
        f = float(m[1])
        return (int(f) if f.is_integer() else f), True
    return None, False


def _date(v):
    """Excel date cell or 'YYYY-MM-DD' text → 'YYYY-MM-DD', else None. openpyxl gives naive datetimes: no zone shift."""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, str) and ISO_RE.match(v.strip()):
        try:
            return datetime.date.fromisoformat(v.strip()).isoformat()
        except ValueError:
            return None
    return None


def parse_template(wb):
    """→ (berths, vessels, rows, issues, n_sheets, n_cells) already in ParsedWorkbook shape."""
    issues, berths, vessels, rows = [], [], [], []
    n_sheets = n_cells = 0
    berth_names, vessel_keys = set(), {}

    def issue(code, severity, sheet, cell, message, row=None):
        issues.append({"code": code, "severity": severity, "sheet": sheet, "cell": cell, "message": message, "row": row})

    def records(name):
        ws = _sheet(wb, name)
        if ws is None:
            return
        nonlocal n_sheets, n_cells
        n_sheets += 1
        header = [str(v).strip().lower() if v is not None else "" for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())]
        want = HEADERS[name]
        if header[:len(want)] != want:
            issue("TEMPLATE_BAD_HEADER", "error", ws.title, "A1",
                  f"Expected the header row to be: {' | '.join(h.title() for h in want)}. This sheet was skipped.")
            return
        for r, values in enumerate(ws.iter_rows(min_row=2, max_col=len(want), values_only=True), start=2):
            values = list(values) + [None] * (len(want) - len(values))
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
                continue
            if isinstance(values[0], str) and values[0].strip().lower().startswith("e.g."):
                continue
            n_cells += sum(v is not None for v in values)
            yield ws.title, r, dict(zip(want, values))

    def ref(r, col_index):
        return f"{get_column_letter(col_index + 1)}{r}"

    for sheet, r, v in records("berths") or ():
        name = _text(v["name"])
        kind = (_text(v["kind"]) or "").lower()
        length, length_ok = _feet(v["length (ft)"])
        if not name:
            issue("INVALID_VALUE", "error", sheet, ref(r, 0), "Berth name is empty; row skipped.")
            continue
        if kind not in ("berth", "section"):
            issue("INVALID_VALUE", "error", sheet, ref(r, 1), f"Kind must be \"berth\" or \"section\" (got \"{v['kind']}\"); row skipped.")
            continue
        if not length_ok or (kind == "berth" and length is None) or (kind == "section" and length is not None):
            need = "a length in feet" if kind == "berth" else "no length"
            issue("INVALID_VALUE", "error", sheet, ref(r, 2), f"A {kind} needs {need}; row skipped.")
            continue
        if name.lower() in berth_names:
            issue("DUPLICATE_NAME", "warning", sheet, ref(r, 0), f"Berth \"{name}\" appears twice; the first one was kept.")
            continue
        order = v["order"] if isinstance(v["order"], int) else len(berths) + 1
        berth_names.add(name.lower())
        berths.append({"name": name, "kind": kind, "lengthFt": length, "sortOrder": order})

    for sheet, r, v in records("vessels") or ():
        name = _text(v["name"])
        length, length_ok = _feet(v["length (ft)"])
        draft, draft_ok = _feet(v["draft (ft)"])
        if not name:
            issue("INVALID_VALUE", "error", sheet, ref(r, 0), "Vessel name is empty; row skipped.")
            continue
        if not length_ok or not draft_ok:
            issue("INVALID_VALUE", "error", sheet, ref(r, 1 if not length_ok else 2), f"\"{name}\": lengths must be positive numbers of feet; row skipped.")
            continue
        key = vessel_key(name)
        if key in vessel_keys:
            issue("DUPLICATE_NAME", "warning", sheet, ref(r, 0), f"Vessel \"{name}\" appears twice; the first one was kept.")
            continue
        vessel_keys[key] = len(vessels)
        vessels.append({"name": vessel_display(name) if re.match(r"^\S+/\S+\s", name) else name,
                        "lengthFt": length, "draftFt": draft, "operator": _text(v["operator"]), "notes": _text(v["notes"])})

    for sheet, r, v in records("bookings") or ():
        berth, typ, title = _text(v["berth"]), (_text(v["type"]) or "").lower(), _text(v["vessel / title"])
        start, end = _date(v["start"]), _date(v["end"])
        bad = None
        if typ not in ("vessel", "event", "closure"):
            bad = (1, f"Type must be vessel, event or closure (got \"{v['type']}\")")
        elif not title:
            bad = (2, "Vessel / Title is empty")
        elif not start:
            bad = (3, "Start must be a date (a date cell or YYYY-MM-DD)")
        elif not end:
            bad = (4, "End must be a date (a date cell or YYYY-MM-DD)")
        if bad:
            issue("INVALID_VALUE", "error", sheet, ref(r, bad[0]), f"{bad[1]}; row skipped.")
            continue
        if typ == "vessel":
            key = vessel_key(title)
            if key not in vessel_keys:            # booked but not listed: add it with an unknown length
                vessel_keys[key] = len(vessels)
                vessels.append({"name": title, "lengthFt": None, "draftFt": None, "operator": None, "notes": None})
            title = vessels[vessel_keys[key]]["name"]
        row = {"berthLabel": berth, "occupantType": typ, "title": title, "startDate": start, "endDate": end,
               "notes": _text(v["notes"]), "sheet": sheet, "cell": ref(r, 0), "classifiedBy": "template"}
        if berth is None:
            issue("NO_BERTH", "error", sheet, ref(r, 0), "Berth is empty: choose one to create this booking.", row)
        else:
            rows.append(row)

    return berths, vessels, rows, issues, n_sheets, n_cells
