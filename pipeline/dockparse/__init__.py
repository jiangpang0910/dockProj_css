"""dockparse — the Python parse pipeline. `run(bytes)` → a ParsedWorkbook dict (shared/pipeline.ts, version 1).

    bytes ─► open (.xlsx, or a CSV as a one-sheet workbook)
          ─► sort every sheet by what it IS, not what it's called:
                calendar grid (month headers + day numbers)  ─► grid.py
                Science / Yachts registry                     ─► registry.py
                exactly our template (name + header row)      ─► template.py
                anything else                                 ─► table.py (free-form table: header found, columns
                                                                 scored by wording + values, records cleaned)
                still unreadable, model enabled               ─► model.py proposes a layout, table.py reads with it
          ─► classify occupants (regex, then model.py for leftovers), normalise names, attach lengths, merge
             touching day ranges ─► ParsedWorkbook. format: "legacy_grid" | "template" | "table" | None.

It answers "what does the file say?" only. It never decides whether a booking is allowed: overlaps, fit and
double-berthing are judged by rules.ts in the TS importer, along with the planning window. Every non-empty cell
gets a disposition in the ledger (audit.py); what can't be read is an issue that says which cell and why.
"""
import csv
import io
import warnings

import openpyxl

from . import model
from .classify import classify, vessel_display, vessel_key
from .extras import read_tours, read_usage
from .grid import is_grid_sheet, parse_grid
from .registry import read_registry
from .table import FIELDS, read_tables
from .template import HEADERS as TEMPLATE_HEADERS, parse_template, template_sheets

VERSION = 1
REGISTRY_SHEETS = ("Science", "Yachts")
# the legacy workbook's extra tabs: known, and known not to hold bookings
REFERENCE_SHEETS = {"Tours": "Tours log (read by extras.py)",
                    "8YR Dock Summary": "usage summary (read by extras.py)"}
MODEL_LAYOUT_MIN_ROWS, MODEL_LAYOUT_MIN_COLS = 3, 2


def run(data: bytes, use_model: bool = True) -> dict:
    wb = _open(data)
    if wb is None:
        return _unknown("Couldn't open the file: it is neither an .xlsx workbook nor a CSV text file.")
    return _parse(wb, use_model, None)


def _open(data: bytes):
    """.xlsx → workbook; CSV text → a one-sheet workbook (values stay text; the readers parse them); else None."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")      # openpyxl warns about data validations / extensions it drops
        try:
            return openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        except Exception:
            pass
    if b"\0" in data[:4096]:
        return None
    for enc in ("utf-8-sig", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        return None
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect)]
    if not any(len([c for c in r if c.strip()]) >= 2 for r in rows):
        return None
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CSV"
    for r in rows:
        ws.append([c.strip() or None for c in r])
    return wb


def _parse(wb, use_model: bool, ledger) -> dict:
    """The whole pipeline over an open workbook. ledger (audit.py): {(sheet, cell): what happened to it}."""
    grid_ws = [ws for ws in wb.worksheets if is_grid_sheet(ws)]
    tmpl = template_sheets(wb)
    other = [ws for ws in wb.worksheets
             if ws not in grid_ws and ws.title not in tmpl and ws.title not in REGISTRY_SHEETS and ws.title not in REFERENCE_SHEETS]
    tours, usage = read_tours(wb, ledger), read_usage(wb, ledger)
    if ledger is not None:
        for ws in wb.worksheets:
            what = ("our template sheet" if ws.title in tmpl else None) or REFERENCE_SHEETS.get(ws.title)
            if what:
                for row in ws.iter_rows():
                    for c in row:
                        if c.value is not None:
                            ledger.setdefault((ws.title, c.coordinate), what)

    berths, vessels, rows, issues = [], {}, [], []
    n_sheets = n_cells = calls = 0
    if grid_ws:
        g = _run_grid(wb, use_model, ledger)
        berths, vessels, rows, issues = g["berths"], g["vessels"], g["rows"], g["issues"]
        n_sheets, n_cells, calls = g["n_sheets"], g["n_cells"], g["calls"]
    if tmpl:
        tb, tv, tr, ti, ts, tc = parse_template(wb, only=tmpl)
        _merge_berths(berths, tb)
        _merge_vessels(vessels, {vessel_key(v["name"]): v for v in tv})
        rows += tr
        issues += ti
        n_sheets += ts
        n_cells += tc

    read_table = False
    if other:
        t = read_tables(wb, other, known_berths=[b["name"] for b in berths], known_vessels=vessels, ledger=ledger)
        read_table = t["n_sheets"] > 0
        skipped = t["skipped"]
        if use_model and any(_worth_asking(wb[title]) for title, _ in skipped):
            asking = [wb[title] for title, _ in skipped if _worth_asking(wb[title])]
            specs, made = model.propose_layout([_profile(ws) for ws in asking])
            calls += made
            specs = {title: spec for title, spec in specs.items() if title in {ws.title for ws in asking}}
            if specs:
                if ledger is not None:               # the first pass marked these SHEET_SKIPPED; the model pass re-marks
                    for title in specs:
                        for k in [k for k in ledger if k[0] == title]:
                            del ledger[k]
                m = read_tables(wb, [wb[title] for title in specs], known_berths=[b["name"] for b in berths],
                                known_vessels=vessels, ledger=ledger, specs=specs)
                for title in specs:
                    if title not in {s for s, _ in m["skipped"]}:
                        m["issues"].append({"code": "MODEL_LAYOUT", "severity": "info", "sheet": title, "cell": None, "row": None,
                                            "message": f"The column layout of sheet \"{title}\" was proposed by the language model "
                                                       f"because nothing in it matched on its own. Check the mapping and the rows."})
                skipped = [s for s in skipped if s[0] not in specs] + m["skipped"]
                read_table = read_table or m["n_sheets"] > 0
                _absorb(berths, vessels, rows, issues, m)
                n_sheets += m["n_sheets"]
                n_cells += m["n_cells"]
        _absorb(berths, vessels, rows, issues, t)
        n_sheets += t["n_sheets"]
        n_cells += t["n_cells"]
        anything = bool(grid_ws or tmpl or read_table)
        for title, reason in skipped:
            if anything:
                bad_tmpl = TEMPLATE_HEADERS.get(title.strip().lower())
                if bad_tmpl:
                    issues.append({"code": "TEMPLATE_BAD_HEADER", "severity": "error", "sheet": title, "cell": "A1", "row": None,
                                   "message": f"Sheet \"{title}\" is named like our template but its columns couldn't be read "
                                              f"({reason}). Expected: {' | '.join(h.title() for h in bad_tmpl)}."})
                else:
                    issues.append({"code": "SHEET_SKIPPED", "severity": "info", "sheet": title, "cell": None, "row": None,
                                   "message": f"Sheet \"{title}\" was skipped: {reason}."})

    if not (grid_ws or tmpl or read_table):
        why = "; ".join(f"\"{t}\": {r}" for t, r in (skipped if other else [])) or "no sheet holds a schedule"
        return _unknown("Nothing in this file could be read as a schedule, a vessel list or a berth list. "
                        f"Looked at every sheet — {why}. Download the template from the import page, or send a table "
                        "with a header row naming the vessel, berth and dates.")

    _share_namesake_lengths(vessels)
    for r in rows:                                   # every stay shows its vessel the way the vessel list spells it
        if r["occupantType"] == "vessel" and vessel_key(r["title"]) in vessels:
            r["title"] = vessels[vessel_key(r["title"])]["name"]
    fmt = "table" if read_table else "template" if tmpl and not grid_ws else "legacy_grid" if grid_ws else "template"
    rows = _dedupe(rows, issues)
    return _out(fmt, n_sheets, n_cells, calls, berths, sorted(vessels.values(), key=lambda v: v["name"]),
                sorted(rows, key=lambda r: (r["startDate"], r["berthLabel"] or "", r["title"])), issues,
                tours=tours, usage=usage)


def _worth_asking(ws):
    return ws.max_row >= MODEL_LAYOUT_MIN_ROWS and ws.max_column >= MODEL_LAYOUT_MIN_COLS


def _profile(ws, rows=8, cols=12):
    """What the model sees of a sheet: its title and the first non-empty rows, as text."""
    out = []
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40), max_col=cols, values_only=True):
        if any(v is not None for v in r):
            out.append([None if v is None else str(v)[:40] for v in r])
        if len(out) >= rows:
            break
    return {"title": ws.title, "rows": out, "fields": sorted(FIELDS)}


def _absorb(berths, vessels, rows, issues, t):
    _merge_berths(berths, t["berths"])
    _merge_vessels(vessels, t["vessels"])
    rows += t["rows"]
    issues += t["issues"]


def _merge_berths(into, more):
    have = {b["name"].lower() for b in into}
    for b in more:
        if b["name"].lower() not in have:
            into.append(dict(b, sortOrder=len(into) + 1))
            have.add(b["name"].lower())


def _merge_vessels(into, more):
    """By key; a stated length is kept, a larger stated length wins, blanks are filled, and the properly cased
    spelling ("R/V Sea Lion" over a booking cell's "r/v sea lion") is the one shown."""
    for k, v in more.items():
        cur = into.get(k)
        if cur is None:
            into[k] = dict(v)
            continue
        for f in ("lengthFt", "draftFt", "operator", "notes"):
            if cur.get(f) is None and v.get(f) is not None:
                cur[f] = v[f]
        if v.get("lengthFt") and cur.get("lengthFt") and v["lengthFt"] > cur["lengthFt"]:
            cur["lengthFt"] = v["lengthFt"]
        if cur["name"] != v["name"] and v["name"] == vessel_display(v["name"]) and cur["name"] != vessel_display(cur["name"]):
            cur["name"] = v["name"]


def _dedupe(rows, issues):
    """The same stay stated twice (a summary tab repeating a year tab) is one stay; the copy is reported."""
    seen, out = set(), []
    for r in rows:
        k = (r["berthLabel"], r["occupantType"], vessel_key(r["title"]), r["startDate"], r["endDate"])
        if k in seen:
            issues.append({"code": "DUPLICATE_CARRYOVER", "severity": "info", "sheet": r["sheet"], "cell": r["cell"], "row": None,
                           "message": f"\"{r['title']}\" {r['startDate']}–{r['endDate']} on {r['berthLabel'] or 'no berth'} "
                                      f"is listed again here; this copy was skipped."})
            continue
        seen.add(k)
        out.append(r)
    return out


def _run_grid(wb, use_model, ledger=None):
    cells, grid_berths, issues, n_sheets, n_cells = parse_grid(wb, ledger)
    registry = read_registry(wb, ledger)

    # 1. classify: regex first, the model only for what regex can't place
    kinds = {c["text"]: classify(c["text"]) for c in cells}
    unknown = [t for t, k in kinds.items() if k == "unknown"]
    labels, calls = model.label(unknown) if use_model else ({}, 0)

    candidates = []
    for c in cells:
        kind, by = kinds[c["text"]], "regex"
        if kind == "unknown" and c["text"] in labels:
            kind, by = labels[c["text"]], "model"
        berth_name = c["berth"][0] if c["berth"] else None
        if kind == "vessel":
            key = vessel_key(c["text"])
            title = registry[key]["name"] if key in registry else vessel_display(c["text"])
        else:
            title = c["text"]
        row = {"berthLabel": berth_name, "occupantType": kind if kind in ("vessel", "event", "closure") else "event",
               "title": title, "startDate": c["start"].isoformat(), "endDate": c["end"].isoformat(), "notes": c.get("note"),
               "sheet": c["sheet"], "cell": c["cell"], "classifiedBy": by}
        if kind == "note":
            issues.append(_issue("ANNOTATION_SKIPPED", "info", c, f"\"{c['text']}\" reads as an operational note, not an "
                                 f"occupant, so it was skipped. Create it as a booking if it was one.", row))
        elif kind == "unknown":
            issues.append(_issue("UNPARSEABLE_CELL", "warning", c, f"Couldn't tell what \"{c['text']}\" is. Create it as "
                                 f"a booking (it defaults to an event) or dismiss it.", row))
        else:
            candidates.append(row)

    # 2. merge touching/overlapping stays of the same occupant on the same berth (merged cells split by month blocks)
    merged = _merge(candidates)

    # 3. rows need a berth; the rest become NO_BERTH issues carrying the row
    rows = []
    for row in merged:
        if row["berthLabel"] is None:
            issues.append({"code": "NO_BERTH", "severity": "error", "sheet": row["sheet"], "cell": row["cell"],
                           "message": f"\"{row['title']}\" is on an unlabeled row, so its berth is unknown. "
                                      f"Choose a berth to create it, or dismiss it.", "row": row})
        else:
            rows.append(row)
        if row["classifiedBy"] == "model":
            issues.append({"code": "MODEL_CLASSIFIED", "severity": "info", "sheet": row["sheet"], "cell": row["cell"],
                           "message": f"\"{row['title']}\" was classified as {row['occupantType']} by the language model. "
                                      f"Check it.", "row": None})

    # 4. berths from the labels (most common length wins), vessels from the registry + the schedule
    berths = [{"name": name,
               "lengthFt": max(b["lengths"], key=b["lengths"].get) if b["lengths"] else None,
               "sortOrder": b["order"]} for name, b in grid_berths.items()]
    vessels = {k: {"name": v["name"], "lengthFt": v["lengthFt"], "draftFt": v["draftFt"],
                   "operator": v.get("operator"), "notes": v.get("notes")}
               for k, v in registry.items()}
    for row in rows + [i["row"] for i in issues if i["code"] == "NO_BERTH"]:
        if row["occupantType"] == "vessel":
            vessels.setdefault(vessel_key(row["title"]), {"name": row["title"], "lengthFt": None, "draftFt": None,
                                                          "operator": None, "notes": None})
    return {"berths": berths, "vessels": vessels, "rows": rows, "issues": issues, "n_sheets": n_sheets,
            "n_cells": n_cells, "calls": calls}


def _share_namesake_lengths(vessels):
    """Same name after the hull prefix ("S/Y Deep Cove" / "F/V Deep Cove") → the same boat for length purposes.
    Only fills a missing length; a length the file states is never overwritten. If namesakes disagree, the larger
    wins (a fit check on the larger value can never let a too-long vessel through)."""
    known: dict = {}
    for v in vessels.values():
        if v["lengthFt"] is not None:
            body = _name_body(v["name"])
            if body not in known or v["lengthFt"] > known[body]["lengthFt"]:
                known[body] = v
    for v in vessels.values():
        src = known.get(_name_body(v["name"])) if v["lengthFt"] is None else None
        if src:
            v["lengthFt"] = src["lengthFt"]
            taken = f"Length taken from {src['name']} (same name)."
            v["notes"] = f"{taken} {v['notes']}" if v.get("notes") else taken


def _name_body(name):
    """'R/V Deep Cove' → 'DEEP COVE'."""
    return vessel_key(name).partition(" ")[2]


def _merge(rows):
    """Same berth + same occupant + ranges touching or overlapping (≤ 1 day apart) → one stay."""
    import datetime as dt
    groups = {}
    for r in rows:
        groups.setdefault((r["berthLabel"], r["occupantType"], vessel_key(r["title"])), []).append(r)
    out = []
    for grp in groups.values():
        grp.sort(key=lambda r: r["startDate"])
        cur = dict(grp[0])
        for r in grp[1:]:
            next_day = (dt.date.fromisoformat(cur["endDate"]) + dt.timedelta(days=1)).isoformat()
            if r["startDate"] <= next_day:
                cur["endDate"] = max(cur["endDate"], r["endDate"])
                if r["classifiedBy"] == "model":
                    cur["classifiedBy"] = "model"
            else:
                out.append(cur)
                cur = dict(r)
        out.append(cur)
    return sorted(out, key=lambda r: (r["startDate"], r["berthLabel"] or "", r["title"]))


def _issue(code, severity, cell, message, row=None):
    return {"code": code, "severity": severity, "sheet": cell["sheet"], "cell": cell["cell"], "message": message, "row": row}


def _unknown(message):
    return _out(None, 0, 0, 0, [], [], [], [{"code": "UNKNOWN_FORMAT", "severity": "error", "sheet": "", "cell": None,
                                             "message": message, "row": None}])


def _out(fmt, n_sheets, n_cells, calls, berths, vessels, rows, issues, tours=None, usage=None):
    return {"version": VERSION, "format": fmt, "stats": {"sheets": n_sheets, "cells": n_cells, "modelCalls": calls},
            "berths": berths, "vessels": vessels, "rows": rows, "issues": issues,
            "tours": tours or [], "usage": usage or []}
