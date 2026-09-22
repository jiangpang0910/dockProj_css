"""The legacy year-per-sheet grid → raw stays (grew out of backend/parse.py).

Layout: one sheet per year; month blocks start at a row whose column A is a month name; berth rows carry labels like
"North Pier West - 410'"; each day is a column. A STAY is drawn as a coloured bar: the name sits in one cell and the
same fill runs across the stay's days (or the bar is one merged cell). Quirks handled:
  - day columns: every day number near the header votes for where day 1 is (col − n + 1); the majority wins, so a
    header with a few overwritten or shifted numbers is still read right
  - weekday-aligned months: columns before day 1 / after the last day are the neighbouring month's days, and are
    dated as such (a bar there is the same stay seen twice; merging joins them)
  - month-header years can be wrong (the 2010 sheet says "NOVEMBER 2018") → year inferred by sequence
  - sheets 2002–04 open with a copy of the previous December → a (year, month) seen before is dropped
  - text between the month header and the first berth row is header debris, not a stay → HEADER_AREA_TEXT
  - unlabeled rows under a SECTION label are more rows of that section; under a berth they are overflow → NO_BERTH
  - a bar with no name continues the stay that ends the day before on the same berth; if none does → UNLABELED_BAR
"""
import calendar
import collections
import datetime
import re

from openpyxl.utils import get_column_letter

from .classify import FT, WEEKDAY_LETTERS, clean, num

MONTHS = ["JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE", "JULY", "AUGUST", "SEPTEMBER", "OCTOBER",
          "NOVEMBER", "DECEMBER"]
MONTH_RE = re.compile(r"^\s*(%s)\b\s*(\d{4})?\s*$" % "|".join(MONTHS), re.I)
BERTH_RE = re.compile(rf"^\s*(?P<name>.+?)\s+-\s+{FT}\s*$")
SECTION_RE = re.compile(r"^\s*(North Finger Piers|Small craft slips)\b", re.I)
SECTION_NAMES = {"north finger piers": "North Finger Piers", "small craft slips": "Small Craft Slips"}
SPILL_DAYS = 7  # a weekday-aligned block shows at most 6 days of each neighbouring month

# Fills that are page background (white, greys used for weekends/headers), never a stay bar.
NEUTRAL_RGB = {"FFFFFFFF", "00FFFFFF", "FFD9D9D9", "FFF2F2F2", "FFBFBFBF", "FFA6A6A6", "FF808080", "00000000"}
NEUTRAL_INDEXED = {"9", "64", "65", "22", "23", "55"}


def berth_from_label(label):
    """'North Pier West - 410'' → ('North Pier West', 'berth', 410); section labels → (name, 'section', None)."""
    if not isinstance(label, str):
        return None
    if (m := BERTH_RE.match(label)):
        name = clean(m["name"])
        return (name.title() if name.isupper() else name, "berth", num(m[2]))
    if (m := SECTION_RE.match(label)):
        return (SECTION_NAMES[m[1].lower()], "section", None)
    return None


def is_grid_sheet(ws) -> bool:
    """A month name in column A with, within three rows of it, a row of weekday letters or of day numbers: the
    calendar layout, whatever the sheet is called. Day-number rows can be mostly blank in the legacy files (the
    parser votes on what's left), so a few numbers are enough; the weekday row is the surer sign.
    (A table with a Month column has neither.)"""
    for r in range(1, min(ws.max_row, 400) + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and MONTH_RE.match(v):
            for rr in range(max(1, r - 2), min(r + 4, ws.max_row + 1)):   # header rows sit on, above or below the month
                vals = [ws.cell(rr, c).value for c in range(2, min(ws.max_column, 40) + 1)]
                days = sum(1 for x in vals if isinstance(x, int) and not isinstance(x, bool) and 1 <= x <= 31)
                letters = sum(1 for x in vals if isinstance(x, str) and clean(x).upper() in WEEKDAY_LETTERS)
                if days >= 5 or letters >= 15:
                    return True
    return False


def is_grid(wb) -> bool:
    return any(is_grid_sheet(ws) for ws in wb.worksheets)


def fill_of(cell):
    """A hashable colour for a stay bar, or None for no fill / page background."""
    f = cell.fill
    if f is None or f.fill_type in (None, "none"):
        return None
    fg = f.fgColor
    if fg.type == "rgb":
        return None if (fg.rgb or "").upper() in NEUTRAL_RGB else ("rgb", fg.rgb.upper())
    if fg.type == "theme":
        return None if fg.theme in (0, 1) and (fg.tint or 0) >= -0.35 else ("theme", fg.theme, round(fg.tint or 0, 2))
    if fg.type == "indexed":
        return None if str(fg.indexed) in NEUTRAL_INDEXED else ("indexed", str(fg.indexed))
    return None


def _text(v):
    """Cell value → occupant text, or None for blanks, numbers, dates and weekday letters."""
    if v is None or isinstance(v, (int, float, datetime.date)):
        return None
    t = clean(v)
    return None if not t or t.upper() in WEEKDAY_LETTERS else t


def parse_grid(wb, ledger=None):
    """→ (cells, berths, issues, n_sheets, n_cells).

    cells:  [{sheet, cell, label, berth, text, start, end, fill}]  (berth = (name, kind, len) or None; start/end = date)
    berths: {name: {"kind", "lengths": {len: count}, "order"}}
    ledger: optional dict, filled with {(sheet, "B12"): what this parser did with that cell} (see audit.py)
    """
    cells, issues, berths, orphans = [], [], {}, []
    mark = (lambda sh, r, c, what: ledger.setdefault((sh, f"{get_column_letter(c)}{r}"), what)) if ledger is not None \
        else (lambda *_: None)
    seen_blocks = set()
    n_sheets = n_cells = 0

    for ws in wb.worksheets:
        if not is_grid_sheet(ws):
            continue
        n_sheets += 1
        sheet_year = int(ws.title) if re.fullmatch(r"\d{4}", ws.title.strip()) else None
        max_col = ws.max_column

        merged_of = {}
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    merged_of[(r, c)] = rng

        heads = []
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if isinstance(v, str) and (m := MONTH_RE.match(v)):
                heads.append((r, MONTHS.index(m[1].upper()) + 1, int(m[2]) if m[2] else None))
                mark(ws.title, r, 1, "month header")
        for r in range(1, heads[0][0] if heads else ws.max_row + 1):
            for c in range(1, max_col + 1):
                mark(ws.title, r, c, "sheet title")

        prev = None
        for i, (hr, mon, label_year) in enumerate(heads):
            end_row = heads[i + 1][0] - 1 if i + 1 < len(heads) else ws.max_row
            prev_hr = heads[i - 1][0] if i else 0

            # Year: trust the first label only if it's within a year of the sheet; after that, go by sequence.
            if prev is None:
                year = label_year if (label_year and (sheet_year is None or abs(label_year - sheet_year) <= 1)) else sheet_year
                if year is None:                      # neither the sheet name nor the header says which year
                    issues.append(_issue("SHEET_SKIPPED", "info", ws.title, f"A{hr}",
                                         f"Sheet \"{ws.title}\" is a calendar grid but neither its name nor its month "
                                         f"headers say which year it is, so it was skipped. Name the sheet after the year."))
                    break
            else:
                year = prev[0] + (1 if mon < prev[1] else 0)
            prev = (year, mon)
            if label_year and label_year != year:
                issues.append(_issue("HEADER_YEAR_MISMATCH", "info", ws.title, f"A{hr}",
                                     f"Header says {MONTHS[mon - 1].title()} {label_year}; by sequence it is {year}. Used {year}."))

            if (year, mon) in seen_blocks:
                issues.append(_issue("DUPLICATE_CARRYOVER", "info", ws.title, f"A{hr}",
                                     f"{MONTHS[mon - 1].title()} {year} already appeared on an earlier sheet; this copy was skipped."))
                for r in range(hr, end_row + 1):
                    for c in range(1, max_col + 1):
                        mark(ws.title, r, c, "DUPLICATE_CARRYOVER block")
                continue
            seen_blocks.add((year, mon))

            # Where the block's data starts: the first labeled row. Everything from the header to it is header area.
            first_data = next((r for r in range(hr + 1, end_row + 1) if berth_from_label(ws.cell(r, 1).value)), None)
            if first_data is None:
                continue

            # Day 1's column: every day number around the header votes col − n + 1.
            votes = collections.Counter()
            for rr in range(max(prev_hr + 1, hr - 2), first_data):
                for c in range(2, max_col + 1):
                    v = ws.cell(rr, c).value
                    if isinstance(v, int) and not isinstance(v, bool) and 1 <= v <= 31:
                        votes[c - v + 1] += 1
                        mark(ws.title, rr, c, "day number")
            col1 = votes.most_common(1)[0][0] if votes else 2
            first = datetime.date(year, mon, 1)
            last = datetime.date(year, mon, calendar.monthrange(year, mon)[1])

            def date_of(c):
                d = first + datetime.timedelta(days=c - col1)
                return d if first - datetime.timedelta(days=SPILL_DAYS) <= d <= last + datetime.timedelta(days=SPILL_DAYS) else None

            for r in range(hr, first_data):  # header debris: text where day numbers / weekday letters belong
                for c in range(2, max_col + 1):
                    v = ws.cell(r, c).value
                    if isinstance(v, str) and clean(v).upper() in WEEKDAY_LETTERS:
                        mark(ws.title, r, c, "weekday letter")
                    if (t := _text(v)):
                        mark(ws.title, r, c, "HEADER_AREA_TEXT")
                        issues.append(_issue("HEADER_AREA_TEXT", "info", ws.title, f"{get_column_letter(c)}{r}",
                                             f"\"{t}\" sits in the {MONTHS[mon - 1].title()} {year} header rows, where day "
                                             f"numbers and weekday letters go, so it was not read as a stay."))

            section = None
            for r in range(first_data, end_row + 1):
                label = ws.cell(r, 1).value
                label = label if isinstance(label, str) and label.strip() else None
                berth = berth_from_label(label)
                if label:
                    mark(ws.title, r, 1, "berth label" if berth else "row label (not a berth)")
                if berth:
                    b = berths.setdefault(berth[0], {"kind": berth[1], "lengths": {}, "order": len(berths) + 1})
                    if berth[2] is not None:
                        b["lengths"][berth[2]] = b["lengths"].get(berth[2], 0) + 1
                    section = berth if berth[1] == "section" else None
                elif label is None and section:
                    berth = section  # another row of a multi-row section
                for c in range(2, max_col + 1):  # numbers in a stay row (times like 1400) are not stays
                    if isinstance(ws.cell(r, c).value, (int, float)):
                        mark(ws.title, r, c, "number in a stay row (a time?), not read")
                for bar in _name_row(_bars(ws, r, max_col, merged_of)):
                    days = [d for d in (date_of(c) for c in range(bar["c0"], bar["c1"] + 1)) if d]
                    ref = f"{get_column_letter(bar['c'])}{r}"
                    for c in range(bar["c0"], bar["c1"] + 1):
                        mark(ws.title, r, c, ("stay name" if c == bar["c"] else "stay bar") if bar["text"] else "UNLABELED_BAR")
                    if bar["text"]:
                        n_cells += 1
                    if not days:
                        if bar["text"]:
                            issues.append(_issue("OUTSIDE_MONTH_COLUMNS", "warning", ws.title, ref,
                                                 f"\"{bar['text']}\" sits outside the day columns of {MONTHS[mon - 1].title()} "
                                                 f"{year}; its days can't be read."))
                        continue
                    stay = {"sheet": ws.title, "cell": ref, "label": label, "berth": berth, "text": bar["text"],
                            "start": min(days), "end": max(days), "fill": bar["fill"], "note": bar.get("note")}
                    (cells if bar["text"] else orphans).append(stay)

    _attach_orphans(cells, orphans, issues)
    return cells, berths, issues, n_sheets, n_cells


def _bars(ws, r, max_col, merged_of):
    """One row → its bars: [{text, c (name's column), c0, c1, fill}]. A bar is one merged cell, a named cell, or a run
    of cells sharing a non-background fill. A new name inside a run starts a new bar; a name inside a run that began
    unnamed takes the whole run (the bar was drawn first and labelled mid-way)."""
    out, cur = [], None
    for c in range(2, max_col + 1):
        rng = merged_of.get((r, c))
        if rng is not None and rng.min_row != r:
            continue  # a vertical merge's lower rows belong to the row above
        anchor = ws.cell(rng.min_row, rng.min_col) if rng else ws.cell(r, c)
        fill = fill_of(anchor)
        text = _text(anchor.value) if (rng is None or c == rng.min_col) else None
        joins = cur is not None and ((rng is not None and rng is cur["rng"]) or (fill is not None and fill == cur["fill"]))
        if text:
            if joins and cur["text"] is None:
                cur.update(text=text, c=c, c1=max(cur["c1"], rng.max_col if rng else c), rng=rng or cur["rng"])
                continue
            if cur:
                out.append(cur)
            cur = {"text": text, "c": c, "c0": c, "c1": c, "fill": fill, "rng": rng}
            if rng:
                cur["c1"] = rng.max_col
        elif joins:
            cur["c1"] = max(cur["c1"], c)
        else:
            if cur:
                out.append(cur)
            cur = {"text": None, "c": c, "c0": c, "c1": c, "fill": fill, "rng": None} if fill is not None else None
    if cur:
        out.append(cur)
    return out


def _name_row(bars):
    """Give unnamed bars in one row a name when the row itself says whose they are:
      1. touching a named bar (its first cell was coloured differently, or the name sits at one end) → part of it
      2. same colour as the bar(s) of exactly one name in this row → that vessel (the file colour-codes occupants)
    The rest stay unnamed (may still continue last month's stay, see _attach_orphans)."""
    out = []
    for b in bars:  # 1. join touching runs; a run between two names goes to the left one (a continuation)
        if out and b["c0"] == out[-1]["c1"] + 1 and (out[-1]["text"] is None) != (b["text"] is None):
            prev = out[-1]
            if prev["text"] is None:  # unnamed run just before a name: the bar starts earlier
                b = {**b, "c0": prev["c0"]}
                out[-1] = b
            else:
                prev["c1"] = b["c1"]
            continue
        out.append(dict(b))
    by_fill = collections.defaultdict(set)
    for b in out:
        if b["text"] and b["fill"] is not None:
            by_fill[b["fill"]].add(b["text"].upper())
    for b in out:  # 2. colour code
        if b["text"] is None and len(names := by_fill.get(b["fill"], ())) == 1:
            owner = next(x for x in out if x["text"] and x["text"].upper() in names)
            b.update(text=owner["text"], note="Name taken from the bar's colour: the same row has it on this vessel's other bar.")
    return out


def _attach_orphans(cells, orphans, issues):
    """An unnamed bar continues the named stay on the same berth row that ends the day before it (or overlaps it):
    typically a stay carried over from the previous month, whose name was written only in that month."""
    by_berth = collections.defaultdict(list)
    for c in cells:
        by_berth[(c["label"], c["berth"])].append(c)
    for o in sorted(orphans, key=lambda o: o["start"]):
        prior = [c for c in by_berth[(o["label"], o["berth"])]
                 if c["start"] <= o["start"] <= c["end"] + datetime.timedelta(days=1)]
        if prior:
            best = max(prior, key=lambda c: (c["fill"] == o["fill"], c["end"]))
            best["end"] = max(best["end"], o["end"])
        else:
            row = {"berthLabel": o["berth"][0] if o["berth"] else None, "occupantType": "event", "title": "Unlabeled bar",
                   "startDate": o["start"].isoformat(), "endDate": o["end"].isoformat(),
                   "notes": "A coloured bar in the schedule with no name.", "sheet": o["sheet"], "cell": o["cell"],
                   "classifiedBy": "regex"}
            issues.append(_issue("UNLABELED_BAR", "warning", o["sheet"], o["cell"],
                                 f"A coloured bar ({o['start'].isoformat()} to {o['end'].isoformat()}) has no name and "
                                 f"doesn't continue a stay: something may hold the berth. Create it as a booking or "
                                 f"dismiss it.", row))


def _issue(code, severity, sheet, cell, message, row=None):
    return {"code": code, "severity": severity, "sheet": sheet, "cell": cell, "message": message, "row": row}
