"""The legacy workbook's reference tabs, read for what they hold rather than skipped.

Tours: a log of guided visits aboard vessels (date, time, guide, guest, head count, vessel, notes). A tour
doesn't hold a berth, so it is never a booking; it is kept as its own record.
8YR Dock Summary: days of use per berth per year, as the facility counted them. Kept as reference figures the
app can show next to its own count.
"""
import datetime
import re

from openpyxl.utils.datetime import from_excel

TOURS_SHEET = "Tours"
SUMMARY_SHEET = "8YR Dock Summary"
_TIME_RE = re.compile(r"^(\d{1,2}):?(\d{2})$")


def read_tours(wb, ledger=None) -> list:
    """→ [{date, time, guide, guest, people, vessel, notes}] in sheet order; [] when there is no Tours tab."""
    if TOURS_SHEET not in wb.sheetnames:
        return []
    ws = wb[TOURS_SHEET]
    mark = (lambda c, what: ledger.setdefault((ws.title, c.coordinate), what)) if ledger is not None else (lambda *_: None)
    header = None
    out = []
    for row in ws.iter_rows():
        vals = [c.value for c in row]
        texts = [str(v).strip().lower() if v is not None else "" for v in vals]
        if header is None:
            if "date" in texts and any("guide" in t for t in texts):
                header = {t: i for i, t in enumerate(texts) if t}
                for c in row:
                    if c.value is not None:
                        mark(c, "tours header")
            else:
                for c in row:
                    if c.value is not None:
                        mark(c, "tours preamble (kept as-is)")
            continue
        if all(v is None for v in vals):
            continue
        col = lambda *names: next((vals[header[n]] for n in names if n in header and header[n] < len(vals)), None)
        date = _date(col("date"))
        if date is None:
            # a line under a tour with text only in the first column continues that tour's notes
            texts_only = [str(v).strip() for v in vals if isinstance(v, str) and v.strip()]
            if out and len(texts_only) == 1 and isinstance(vals[0], str):
                out[-1]["notes"] = f"{out[-1]['notes']}; {texts_only[0]}" if out[-1]["notes"] else texts_only[0]
                for c in row:
                    if c.value is not None:
                        mark(c, "tour (continued note)")
            else:
                for c in row:
                    if c.value is not None:
                        mark(c, "tours row without a date (skipped)")
            continue
        people = col("people")
        out.append({"date": date, "time": _time(col("time")), "guide": _s(col("guide")), "guest": _s(col("guest")),
                    "people": int(people) if isinstance(people, (int, float)) and not isinstance(people, bool) else None,
                    "vessel": _s(col("dock/ ship", "dock/ship", "vessel", "ship")), "notes": _s(col("notes"))})
        for c in row:
            if c.value is not None:
                mark(c, "tour")
    return out


def read_usage(wb, ledger=None) -> list:
    """→ [{berth, year, days}] from the summary table: first row = title + years, following rows = berth + counts."""
    if SUMMARY_SHEET not in wb.sheetnames:
        return []
    ws = wb[SUMMARY_SHEET]
    mark = (lambda c, what: ledger.setdefault((ws.title, c.coordinate), what)) if ledger is not None else (lambda *_: None)
    years = None
    out = []
    for row in ws.iter_rows():
        vals = [c.value for c in row]
        if all(v is None for v in vals):
            continue
        if years is None:
            years = {i: int(v) for i, v in enumerate(vals) if isinstance(v, (int, float)) and 1900 <= v <= 2100}
            if years:
                for c in row:
                    if c.value is not None:
                        mark(c, "usage summary header")
            continue
        name = vals[0]
        if not isinstance(name, str) or not name.strip():
            continue
        for i, year in years.items():
            v = vals[i] if i < len(vals) else None
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                out.append({"berth": name.strip(), "year": year, "days": int(v)})
        for c in row:
            if c.value is not None:
                mark(c, "usage summary")
    return out


def _date(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 20000 < v < 80000:
        return from_excel(v).date().isoformat()
    if isinstance(v, str):
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d %b %Y", "%b %d, %Y"):
            try:
                return datetime.datetime.strptime(v.strip(), fmt).date().isoformat()
            except ValueError:
                pass
    return None


def _time(v):
    """1530 → "15:30"; 930 → "09:30"; datetime.time → "HH:MM"; other text ("tbd") kept; blank → None."""
    if v is None:
        return None
    if isinstance(v, datetime.time):
        return v.strftime("%H:%M")
    if isinstance(v, datetime.datetime):
        return v.strftime("%H:%M")
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        n = int(v)
        if 0 <= n <= 2359 and n % 100 < 60:
            return f"{n // 100:02d}:{n % 100:02d}"
        return str(n)
    s = str(v).strip()
    m = _TIME_RE.match(s)
    if m and int(m[1]) < 24 and int(m[2]) < 60:
        return f"{int(m[1]):02d}:{m[2]}"
    return s or None


def _s(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None
